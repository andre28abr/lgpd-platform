"""Regressões das melhorias de produto (Parte 5): export do ROPA, reset de senha,
auditoria encadeada e remetente por empresa."""
import io

from sqlalchemy import text

import models
from extensions import db
from services.auditoria import registrar, verificar_cadeia
from services.senha import gerar_token


def test_ropa_exporta_excel_e_pdf(login, client, app):
    login("dpo@acme.com.br")
    r = client.get("/ropa/exportar.xlsx")
    assert r.status_code == 200
    assert "spreadsheetml" in r.headers["Content-Type"]
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(r.data)).active
    assert ws.title == "ROPA"
    assert [c.value for c in ws[4]][:3] == ["Atividade", "Setor", "Titulares"]

    r = client.get("/ropa/relatorio.pdf")
    assert r.status_code == 200
    assert r.headers["Content-Type"].startswith("application/pdf")


def test_reset_de_senha_por_link(client, csrf, login, app):
    with app.app_context():
        u = models.Usuario(empresa_id=1, nome="Reset", email="reset@acme.com.br")
        u.definir_senha("lgpd1234")
        db.session.add(u)
        db.session.commit()
        uid = u.id

    # Pedido do link: mesma resposta para e-mail conhecido e desconhecido.
    client.get("/login")
    for email in ("reset@acme.com.br", "ninguem@acme.com.br"):
        r = client.post("/senha/esqueci", data={"csrf_token": csrf(), "email": email}, follow_redirects=True)
        assert r.status_code == 200
        assert "Se o e-mail estiver cadastrado" in r.get_data(as_text=True)

    with app.app_context():
        token = gerar_token(db.session.get(models.Usuario, uid))

    assert client.get(f"/senha/redefinir/{token}").status_code == 200
    r = client.post(f"/senha/redefinir/{token}", data={
        "csrf_token": csrf(), "senha": "novaSenha123", "confirmacao": "novaSenha123",
    })
    assert r.status_code in (302, 303) and "/login" in r.headers["Location"]

    # O mesmo link não vale duas vezes (o hash da senha mudou).
    r = client.get(f"/senha/redefinir/{token}")
    assert r.status_code in (302, 303) and "/senha/esqueci" in r.headers["Location"]

    # A senha nova funciona.
    r = login("reset@acme.com.br", "novaSenha123")
    assert r.status_code in (302, 303) and "/login" not in r.headers["Location"]


def test_token_invalido_e_rejeitado(client):
    r = client.get("/senha/redefinir/token-qualquer")
    assert r.status_code in (302, 303) and "/senha/esqueci" in r.headers["Location"]


def test_auditoria_encadeada_detecta_adulteracao(app):
    with app.app_context():
        a = registrar("cadeia_teste_1", "primeiro", empresa_id=1, commit=True)
        b = registrar("cadeia_teste_2", "segundo", empresa_id=1, commit=True)
        assert a.hash and b.hash and a.hash != b.hash
        ok, total, _ = verificar_cadeia()
        assert ok and total >= 2

        # Adultera um registro "por fora" (como alguém com acesso direto ao banco faria).
        db.session.execute(text("UPDATE audit_logs SET detalhe = 'alterado' WHERE id = :id"), {"id": a.id})
        db.session.commit()
        ok, _, quebrado_em = verificar_cadeia()
        assert not ok and quebrado_em == a.id

        # Restaura para não afetar outros testes.
        db.session.execute(text("UPDATE audit_logs SET detalhe = 'primeiro' WHERE id = :id"), {"id": a.id})
        db.session.commit()
        assert verificar_cadeia()[0]


def test_remetente_por_empresa(login, client, csrf, app):
    login("dpo@acme.com.br")
    r = client.post("/admin/configuracoes", data={
        "csrf_token": csrf(), "email_remetente": "Privacidade@Acme.com.br",
    })
    assert r.status_code in (302, 303)
    with app.app_context():
        empresa = models.Empresa.query.filter_by(slug="acme").first()
        assert empresa.email_remetente == "privacidade@acme.com.br"
        empresa.email_remetente = None
        db.session.commit()

    r = client.post("/admin/configuracoes", data={"csrf_token": csrf(), "email_remetente": "sem-arroba"},
                    follow_redirects=True)
    assert "remetente válido" in r.get_data(as_text=True)
