"""Importação do ROPA a partir de planilha (xlsx ou csv) — o inverso da exportação.

O DPO costuma chegar com o inventário numa planilha. As colunas seguem o cabeçalho do
export/modelo; nomes são reconhecidos sem acento e sem distinção de maiúsculas.
"""
import csv
import io
import unicodedata

import models
from extensions import db

CABECALHO = ["Atividade", "Setor", "Titulares", "Categorias de dados", "Finalidade",
             "Base legal", "Retenção", "Compartilhamento"]
_CAMPOS = ["atividade", "setor", "titulares", "categorias_dados", "finalidade",
           "base_legal", "retencao", "compartilhamento"]
_SINONIMOS = {
    "atividade": "atividade", "atividade de tratamento": "atividade", "tratamento": "atividade",
    "setor": "setor", "area": "setor", "departamento": "setor",
    "titulares": "titulares", "titular": "titulares",
    "categorias de dados": "categorias_dados", "categorias": "categorias_dados", "dados": "categorias_dados",
    "finalidade": "finalidade", "finalidades": "finalidade",
    "base legal": "base_legal", "hipotese legal": "base_legal", "base": "base_legal",
    "retencao": "retencao", "prazo de retencao": "retencao", "guarda": "retencao",
    "compartilhamento": "compartilhamento", "compartilhado com": "compartilhamento",
}


def _norm(texto) -> str:
    base = unicodedata.normalize("NFKD", str(texto or "")).encode("ascii", "ignore").decode()
    return " ".join(base.lower().split())


def modelo_xlsx() -> io.BytesIO:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "ROPA"
    ws.append(CABECALHO)
    ws.append(["Folha de pagamento", "Recursos Humanos", "Colaboradores", "Nome, CPF, conta bancária",
               "Processar a folha", "Obrigação legal/regulatória (Art. 7º, II)",
               "Durante o vínculo + prazos legais", "Banco, contabilidade"])
    for coluna, largura in zip("ABCDEFGH", (32, 22, 22, 36, 36, 34, 26, 30), strict=True):
        ws.column_dimensions[coluna].width = largura
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def ler_planilha(arquivo, nome: str) -> list[dict]:
    """Lê xlsx ou csv e devolve linhas como dicts com as chaves de ``_CAMPOS``.
    Levanta ValueError com mensagem para o usuário."""
    ext = nome.rsplit(".", 1)[-1].lower() if "." in nome else ""
    if ext == "xlsx":
        from openpyxl import load_workbook

        ws = load_workbook(io.BytesIO(arquivo.read()), read_only=True, data_only=True).active
        linhas = [[c for c in row] for row in ws.iter_rows(values_only=True)]
    elif ext == "csv":
        texto = arquivo.read().decode("utf-8-sig", errors="replace")
        delimitador = ";" if texto.count(";") > texto.count(",") else ","
        linhas = list(csv.reader(io.StringIO(texto), delimiter=delimitador))
    else:
        raise ValueError("Envie uma planilha .xlsx ou .csv.")

    linhas = [linha for linha in linhas if any(str(c or "").strip() for c in linha)]
    if not linhas:
        raise ValueError("A planilha está vazia.")
    cabecalho = [_SINONIMOS.get(_norm(c)) for c in linhas[0]]
    if "atividade" not in cabecalho:
        raise ValueError("Não encontrei a coluna 'Atividade'. Use o modelo (Baixar modelo) como referência.")

    registros = []
    for linha in linhas[1:]:
        reg = {campo: "" for campo in _CAMPOS}
        for campo, valor in zip(cabecalho, linha, strict=False):
            if campo:
                reg[campo] = str(valor or "").strip()
        if reg["atividade"]:
            registros.append(reg)
    return registros


def _base_legal(valor):
    """Aceita o código (CONSENTIMENTO) ou o rótulo ('Consentimento (Art. 7º, I)'), sem acento."""
    v = _norm(valor)
    if not v:
        return None
    for code, label in models.BASES_LEGAIS:
        if v in (code.lower(), _norm(label)) or v == _norm(label.split("(")[0]):
            return code
    return None


def importar(empresa_id: int, registros: list[dict]) -> tuple[int, list[str]]:
    """Cria os RopaRegistro. Retorna (criados, avisos)."""
    setores = {_norm(s.nome): s.id for s in models.Setor.query.filter_by(empresa_id=empresa_id).all()}
    criados, avisos = 0, []
    for i, reg in enumerate(registros, start=2):
        setor_id = setores.get(_norm(reg["setor"])) if reg["setor"] else None
        if reg["setor"] and setor_id is None:
            avisos.append(f"linha {i}: setor '{reg['setor']}' não existe — registro salvo como 'Toda a empresa'")
        base = _base_legal(reg["base_legal"])
        if reg["base_legal"] and base is None:
            avisos.append(f"linha {i}: base legal '{reg['base_legal']}' não reconhecida — deixada em branco")
        db.session.add(models.RopaRegistro(
            empresa_id=empresa_id, setor_id=setor_id, atividade=reg["atividade"][:200],
            titulares=reg["titulares"][:255], categorias_dados=reg["categorias_dados"],
            finalidade=reg["finalidade"], base_legal=base, retencao=reg["retencao"][:255],
            compartilhamento=reg["compartilhamento"],
        ))
        criados += 1
    return criados, avisos
